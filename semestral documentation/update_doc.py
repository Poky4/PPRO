import docx
from docx.shared import Pt, RGBColor

def main():
    template_path = "/tmp/semestral documentation/PPRO-dokumentace-sablona-v1.docx"
    doc = docx.Document(template_path)

    def format_cell(cell, text, bold=False, italic=False, font_size=9, color=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(font_size)
        if color:
            r.font.color.rgb = color
        else:
            r.font.color.rgb = RGBColor(0, 0, 0)
        return r

    # --- Table 0: Metadata ---
    t0 = doc.tables[0]
    format_cell(t0.rows[0].cells[1], "Rezervace v ordinaci", bold=False, italic=False, font_size=10)
    format_cell(t0.rows[1].cells[1], "Zadání A: Rezervace v ordinaci", bold=False, italic=False, font_size=10)
    format_cell(t0.rows[2].cells[1], "[Doplňte jméno a příjmení]", bold=False, italic=True, font_size=10, color=RGBColor(118, 118, 118))
    format_cell(t0.rows[3].cells[1], "https://github.com/Poky4/PPRO", bold=False, italic=False, font_size=10)
    format_cell(t0.rows[4].cells[1], "19. 10. 2026", bold=False, italic=False, font_size=10)

    # --- Section 1: Popis řešení ---
    # Replace paragraph 5 (the italic guideline) with comprehensive solution description
    p5 = doc.paragraphs[5]
    p5.text = ""
    r_a = p5.add_run(
        "Aplikace slouží jako informační a rezervační systém pro soukromou lékařskou praxi „Ordinace Na Vyhlídce“ "
        "(vedoucí lékařka MUDr. Jana Hrubá, celkem 3 lékaři na částečný úvazek a personál recepce). "
        "Hlavním cílem systému je plnohodnotná náhrada dosavadního papírového diáře na recepci a odstranění "
        "kritických provozních chyb: zamezení vzniku dvojích rezervací na tentýž termín u jednoho lékaře, "
        "automatické respektování specifických ordinačních hodin lékařů (včetně dnů, kdy lékař ordinuje na "
        "externí adrese), rychlé vyhledávání pacientů podle příjmení a telefonního čísla při telefonickém hovoru "
        "a automatizace měsíčního vykazování provedených zdravotních výkonů a bodů pro pojišťovny."
    )
    r_a.italic = False
    r_a.font.size = Pt(11)
    r_a.font.color.rgb = RGBColor(0, 0, 0)
    p5.paragraph_format.space_after = Pt(6)

    # Paragraph 6: Three-tier architecture breakdown
    p6 = doc.paragraphs[6]
    p6.text = ""
    r_b1 = p6.add_run("Systém je navržen v souladu se zásadami čisté třívrstvé architektury se striktně jednosměrným tokem závislostí:\n")
    r_b1.italic = False
    r_b1.font.size = Pt(11)
    r_b1.font.color.rgb = RGBColor(0, 0, 0)

    r_b2 = p6.add_run(
        "• Prezentační vrstva: Poskytuje REST API rozhraní postavené na asynchronním frameworku FastAPI "
        "s automatickou interaktivní dokumentací Swagger UI (/docs) a odlehčeným webovým uživatelským rozhraním (SPA) "
        "optimalizovaným pro ergonomickou práci personálu recepce.\n"
        "• Aplikační a byznys vrstva (Services): Zapouzdřuje doménové služby a veškerá byznys pravidla. "
        "Garantuje klíčové invarianty ordinace – detekci a zákaz kolizí v rezervacích, validaci ordinačních hodin, "
        "správu stavů termínů (plánováno, dokončeno, zrušeno) a kalkulaci bodového ohodnocení návštěv podle číselníku výkonů.\n"
        "• Datová a infrastrukturní vrstva: Odpovídá za perzistenci dat v relační databázi PostgreSQL 16. "
        "Mapování doménových entit zajišťuje SQLAlchemy 2.0 ORM a veškeré změny databázového schématu jsou řízeny "
        "verzovanými migracemi přes nástroj Alembic."
    )
    r_b2.italic = False
    r_b2.font.size = Pt(11)
    r_b2.font.color.rgb = RGBColor(0, 0, 0)
    p6.paragraph_format.space_after = Pt(6)

    # Paragraph 7: Protocols and deployment
    p7 = doc.paragraphs[7]
    p7.text = ""
    r_c = p7.add_run(
        "Klientská část komunikuje se serverem prostřednictvím standardizovaných HTTP/JSON požadavků "
        "zabezpečených pomocí bezstavových JWT Bearer tokenů (OAuth2). Celý systém včetně databáze je kompletně "
        "zkontejnerizován v Dockeru a spustitelný jediným příkazem „docker compose up“ bez nutnosti jakékoliv manuální konfigurace prostředí."
    )
    r_c.italic = False
    r_c.font.size = Pt(11)
    r_c.font.color.rgb = RGBColor(0, 0, 0)
    p7.paragraph_format.space_after = Pt(12)

    # Clean unused empty paragraphs 8-11
    for p_idx in [8, 9, 10, 11]:
        doc.paragraphs[p_idx].text = ""

    # --- Table 1: Rozhodnutí ---
    t1 = doc.tables[1]
    decisions = [
        (
            "R1",
            "Třívrstvá architektura se striktně jednosměrným tokem závislostí (Prezentace → Aplikační vrstva → Databáze)",
            "Důsledné oddělení zodpovědností, nezávislost byznys pravidel na způsobu prezentace a databázi, vysoká testovatelnost doménové logiky a splnění povinného minima PPRO.",
            "12. 10. 2026",
        ),
        (
            "R2",
            "Backend postavený na jazyce Python 3.12 a frameworku FastAPI",
            "Moderní asynchronní webový framework s nativní deklarativní validací vstupů (Pydantic), automaticky generovanou dokumentací OpenAPI/Swagger a vysokou rychlostí vývoje.",
            "12. 10. 2026",
        ),
        (
            "R3",
            "Relační databáze PostgreSQL 16 provozovaná v Docker kontejneru",
            "Robustní databázový systém s precizní podporou časových razítek a transakčních zámků; běh v Dockeru zajišťuje identické běhové prostředí pro vývoj i hodnocení.",
            "12. 10. 2026",
        ),
        (
            "R4",
            "Verzovaná správa databázového schématu nástrojem Alembic (SQLAlchemy 2.0)",
            "Požadavek na verzované migrace; umožňuje deterministické a opakovatelné povyšování schématu bez ad-hoc ručních zásahů do databáze.",
            "12. 10. 2026",
        ),
        (
            "R5",
            "Autentizace pomocí OAuth2 Password Flow s JWT Bearer tokeny a hešováním bcrypt",
            "Bezstavové zabezpečení REST API, bezpečné ukládání přístupových hesel a připravenost pro budoucí rozlišení rolí (Admin, Lékař, Recepce).",
            "14. 10. 2026",
        ),
        (
            "R6",
            "Striktní zákaz kolidujících rezervací a validace ordinačních hodin v byznys vrstvě i DB",
            "Klíčový invariant zadání – zamezení vzniku dvou překrývajících se rezervací u téhož lékaře a vyloučení rezervací mimo ordinační dobu či na jiné adrese.",
            "14. 10. 2026",
        ),
        (
            "R7",
            "Stavový životní cyklus rezervací a zachování historie (soft delete / status model)",
            "Zrušená rezervace uvolní termín v kalendáři pro nového pacienta, ale v databázi zůstává zachována pro evidenci nespolehlivých pacientů a auditní dohledatelnost.",
            "15. 10. 2026",
        ),
        (
            "R8",
            "Automatizované testování pomocí pytest s oddělenou in-memory testovací databází",
            "Izolované, rychlé a deterministické ověřování byznys logiky, kolizních stavů i API endpointů bez závislosti na běžícím externím serveru.",
            "16. 10. 2026",
        ),
        (
            "R9",
            "Kompletní kontejnerizace aplikace a databáze s jednopříkazovým startem (docker compose up)",
            "Splnění povinného minima č. 4 – nulové ruční kroky při spouštění novým uživatelem, automatická inicializace spojení, migrací a seed dat při startu.",
            "16. 10. 2026",
        ),
        (
            "R10",
            "Redesign UI na střídmý lékařský styl a odebrání Swagger odkazu z klientské lišty",
            "Reakce na požadavek klientky (Z1, Z2); odstranění generického AI/SaaS stylu (fialové záře, kulatý avatar) a posílení důvěryhodnosti pro zdravotnický personál.",
            "19. 10. 2026",
        ),
    ]

    for idx, dec in enumerate(decisions, start=1):
        row = t1.rows[idx]
        format_cell(row.cells[0], dec[0], bold=False, italic=False, font_size=9)
        format_cell(row.cells[1], dec[1], bold=False, italic=False, font_size=9)
        format_cell(row.cells[2], dec[2], bold=False, italic=False, font_size=9)
        format_cell(row.cells[3], dec[3], bold=False, italic=False, font_size=9)

    for row in t1.rows[len(decisions) + 1:]:
        for c in row.cells:
            format_cell(c, "", bold=False, italic=False, font_size=9)

    # --- Table 2: Změny požadované klientem ---
    t2 = doc.tables[2]
    client_changes = [
        (
            "Z1",
            "19. 10. 2026",
            "„Nelíbí se mi, že aplikace působí jako vygenerovaná umělou inteligencí. Chci pryč tu generickou ikonku, upravit název na naši ordinaci a zbavit se toho kulatého uživatelského profilu.“",
            "Přepracován vizuální styl klientského rozhraní – fialové záře nahrazeny střídmou lékařskou modří a šedými tóny, generická ikonka nahrazena symbolem lékařského kříže, doplněna identifikace praxe MUDr. Jany Hrubé a kulatý avatar nahrazen hranatým stavovým štítkem personálu. Viz R10.",
            "hotovo",
        ),
        (
            "Z2",
            "19. 10. 2026",
            "„V horní liště je tlačítko 'Swagger API'. Sestřička ani lékaři nevědí, co to je, v hotovém systému pro ordinaci to nemá co dělat.“",
            "Vývojářský odkaz na Swagger dokumentaci byl odstraněn z uživatelského rozhraní; technická dokumentace zůstává dostupná pouze vývojářům a hodnotiteli na přímé adrese /docs. Viz R10.",
            "hotovo",
        ),
    ]

    for idx, change in enumerate(client_changes, start=1):
        row = t2.rows[idx]
        format_cell(row.cells[0], change[0], bold=False, italic=False, font_size=9)
        format_cell(row.cells[1], change[1], bold=False, italic=False, font_size=9)
        format_cell(row.cells[2], change[2], bold=False, italic=False, font_size=9)
        format_cell(row.cells[3], change[3], bold=False, italic=False, font_size=9)
        format_cell(row.cells[4], change[4], bold=False, italic=False, font_size=9)

    for row in t2.rows[len(client_changes) + 1:]:
        for c in row.cells:
            format_cell(c, "", bold=False, italic=False, font_size=9)

    out_sablona = "/tmp/semestral documentation/PPRO-dokumentace-sablona-v1.docx"
    out_spec = "/tmp/semestral documentation/PPRO-dokumentace-Rezervace-v-ordinaci.docx"
    doc.save(out_sablona)
    doc.save(out_spec)
    print("Dokumenty úspěšně aktualizovány a uloženy!")

if __name__ == "__main__":
    main()
